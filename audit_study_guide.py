from docx import Document
from pathlib import Path
from zipfile import ZipFile

path = Path('output/VAE_and_GAN_Simple_Explanation_15_Mark_Study_Guide.docx')
doc = Document(path)
headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
required = [
    '1. KL Divergence and Its Direction',
    '2. Self Information Entropy Cross Entropy and KL',
    '3. VAE Architecture and Encoder',
    '4. VAE Decoder ELBO and Loss',
    '5. Reparameterization Trick',
    '6. Complete VAE Working in Simple Words',
    '7. Limitations of VAEs and Motivation for GANs',
    '8. GAN Architecture and Alternating Training',
    '9. GAN Objectives and Minimax Game',
    '10. GAN Convergence and Nash Equilibrium',
]
missing = [x for x in required if x not in headings]
assert not missing, f'Missing headings: {missing}'
assert len(doc.tables) == 1, 'Expected exactly one comparison table'
table = doc.tables[0]
assert len(table.rows) == 5 and len(table.columns) == 3, 'Comparison table shape is wrong'
text = '\n'.join(p.text for p in doc.paragraphs)
for phrase in ['D_KL(P || Q)', 'z = mu + sigma', 'min_G max_D', 'p_g=p_data']:
    assert phrase in text, f'Missing key content: {phrase}'
for phrase in ['Lecture graph explained', 'How to read it:', 'One Minute Memory Stories']:
    assert phrase in text, f'Missing memorisation support: {phrase}'
assert text.count('15 Mark Answer Points') == 10, 'Each topic must include a 15-mark answer section'
assert text.count('Full process flowchart') == 10, 'Each topic must include a full process flowchart'
assert text.count('Simple model answer to learn') == 10, 'Each topic must include a plain-language model answer'
assert text.count('What each part does in simple words') == 10, 'Each topic must include simple component explanations'
with ZipFile(path) as archive:
    images = [name for name in archive.namelist() if name.startswith('word/media/')]
assert len(images) == 16, f'Expected six lecture visuals plus ten process flowcharts; found {len(images)}'
print(f'OK: {len(doc.paragraphs)} paragraphs, {len(headings)} headings, {len(doc.tables)} table')
print(f'Embedded visuals: {len(images)}')
print(f'File bytes: {path.stat().st_size}')
