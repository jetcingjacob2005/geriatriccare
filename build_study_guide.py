from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path("output") / "VAE_and_GAN_Simple_Explanation_15_Mark_Study_Guide.docx"
OUT.parent.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.72); sec.right_margin = Inches(0.72)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.12
for s, size in [('Title', 23), ('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11)]:
    styles[s].font.name = 'Aptos Display' if s != 'Normal' else 'Aptos'
    styles[s]._element.rPr.rFonts.set(qn('w:ascii'), styles[s].font.name)
    styles[s].font.size = Pt(size); styles[s].font.color.rgb = RGBColor(0, 0, 0)
styles['Heading 1'].font.bold = True; styles['Heading 1'].paragraph_format.space_before = Pt(16); styles['Heading 1'].paragraph_format.space_after = Pt(7)
styles['Heading 2'].font.bold = True; styles['Heading 2'].paragraph_format.space_before = Pt(11); styles['Heading 2'].paragraph_format.space_after = Pt(4)

def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('Page ')
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld)

footer = sec.footer.paragraphs[0]
footer.text = 'VAE and GAN Study Guide | Exam Revision'
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.runs[0].font.size = Pt(8)

def add_title(text, subtitle=None):
    p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    if subtitle:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitle); r.bold = True; r.font.size = Pt(11)

def add_para(text='', bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
    p.add_run(text)
    return p

def bullets(items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item)

def flow(lines):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(7)
    r = p.add_run('\n'.join(lines)); r.font.name = 'Consolas'; r._element.rPr.rFonts.set(qn('w:ascii'), 'Consolas'); r.font.size = Pt(9)

def memory(items):
    doc.add_paragraph('Memory points for a 15-mark answer', style='Heading 3')
    bullets(items)

def visual(filename, caption, explanation):
    """Place a lecture visual next to a short, exam-friendly explanation."""
    image_path = Path('tmp_source_graphs') / filename
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    doc.add_paragraph('Lecture graph explained', style='Heading 2')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(6.35))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption); r.italic = True; r.font.size = Pt(9)
    add_para(explanation, 'How to read it: ')

def topic_flowchart(number):
    chart = Path('tmp_topic_flowcharts') / f'topic_{number}_flowchart.png'
    if not chart.exists():
        raise FileNotFoundError(chart)
    doc.add_paragraph('Full process flowchart', style='Heading 2')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(chart), width=Inches(6.35))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Use this chart to recall the complete order of the answer before writing the formula.'); r.italic = True; r.font.size = Pt(9)

fifteen_mark_points = {
    '1': [
        'Begin by defining P(x) as the true/reference distribution and Q(x) as the learned/model distribution. Explain that KL measures distribution mismatch, not geometric distance between two points.',
        'Write D_KL(P||Q) = sum P(x) log[P(x)/Q(x)]. Explain every part: P weights how important each true-data region is, while the ratio compares true and model probabilities.',
        'Explain non-negativity and asymmetry. KL is zero when P and Q are the same, but reversing the order usually changes the value and therefore the behaviour.',
        'For forward KL, state the key case P(x)>0 and Q(x)=0. The penalty becomes infinite, so Q is forced to cover every region where real data exist.',
        'For reverse KL, state the key case Q(x)>0 where P(x) is very small. The penalty becomes large, so Q avoids unrealistic regions and concentrates on a high-density mode.',
        'Use the two-height-cluster graph as the example and end with the trade-off: forward KL represents all modes but can create in-between samples; reverse KL can be sharp but may miss a valid mode.'
    ],
    '2': [
        'Start with self-information: I(x)=-log2 P(x). It is the surprise of one outcome; rare outcomes carry more information than common outcomes.',
        'Explain why a logarithm is used: for independent events, probabilities multiply but information adds. With base 2, the unit is bits.',
        'Define entropy H(P)=-sum P(x)log2 P(x) as average surprise before the outcome is observed. A fair binary event has maximum uncertainty of one bit.',
        'Define cross-entropy H(P,Q)=-sum P(x)log2 Q(x). It measures the average surprise according to the model Q when reality follows P.',
        'For a one-hot target, show that cross-entropy reduces to -log Q(correct class). This explains why predicting high probability for the true class produces low loss.',
        'Derive KL by subtraction: H(P,Q)-H(P)=sum P(x)log2[P(x)/Q(x)]. Finish by saying KL is the extra uncertainty caused by using Q instead of P.'
    ],
    '3': [
        'Introduce a VAE as an unsupervised generative model: it learns from x without class labels and then can sample new data that resemble the training distribution.',
        'Separate the two spaces. x lives in high-dimensional data space; z lives in a lower-dimensional latent space containing meaningful hidden factors.',
        'Explain the encoder mapping q_phi(z|x). The exact posterior p(z|x) is hard because its denominator p(x) requires integration over every possible z.',
        'State the Gaussian assumption: q_phi(z|x)=N(mu_phi(x), sigma_phi squared(x)). The encoder outputs a mean vector and a log-variance vector, not a single z.',
        'Explain the standard normal prior p(z)=N(0,I). It makes sampling simple and makes nearby points in latent space decode into related outputs.',
        'Conclude with the VAE pathway: encode x into a distribution, sample z, decode z, and later train so the posterior stays close to the prior.'
    ],
    '4': [
        'State the decoder role: it receives z and defines the likelihood p_theta(x|z), meaning how probable an output x is for that latent code.',
        'For continuous data, describe Gaussian decoder outputs such as mean and variance. For binary data, describe Bernoulli probabilities. Link this directly to MSE and BCE.',
        'Explain why several z values can correspond to one x: the encoder produces a distribution. Different samples give slightly different reconstructions that should still resemble the original.',
        'Define the ELBO as E_q[log p_theta(x|z)] - D_KL(q_phi(z|x)||p(z)). Say it is used because log p_theta(x) is hard to compute directly.',
        'Convert maximization to minimization: VAE loss equals negative reconstruction log-likelihood plus KL divergence. Identify the first as accuracy and the second as latent-space regularity.',
        'Conclude that a good VAE must do both jobs at once: create x-hat close to x and keep q(z|x) near the standard normal prior.'
    ],
    '5': [
        'State the problem first: direct random sampling z from N(mu,sigma squared) does not provide a simple differentiable path from z back to the encoder outputs.',
        'Introduce independent noise epsilon sampled from N(0,I). The key fact is that epsilon is random but does not depend on the encoder parameters.',
        'Write the key equation z=mu+sigma elementwise-multiply epsilon. This produces the same Gaussian sample distribution as direct sampling.',
        'Treat epsilon as fixed in one backward pass. Then dz/dmu=1, so a change in mu directly shifts z.',
        'Also show dz/dsigma=epsilon, so changing sigma changes how strongly random noise affects z. Gradients can now reach both encoder outputs.',
        'End with the benefit: the decoder, loss, encoder, and its parameters phi can be trained end-to-end by standard backpropagation.'
    ],
    '6': [
        'Begin with the given input vector and multiply it by encoder weights, add bias, and apply ReLU to form a hidden representation. Mention that ReLU keeps positive values and sets negative values to zero.',
        'Use two separate linear heads to calculate mu and log sigma squared. Do not apply ReLU to these heads because a mean and a log-variance may validly be negative.',
        'Convert log sigma squared with exp, then take its square root to obtain sigma. This prevents the common error of using variance in place of standard deviation.',
        'Draw epsilon from N(0,I) and form z=mu+sigma epsilon. Clearly substitute the supplied epsilon values in the numerical answer.',
        'Pass z through decoder hidden layers and output layer. Choose linear activation for continuous output; choose sigmoid if outputs are binary or normalized to [0,1].',
        'Calculate reconstruction loss and KL loss separately, then add them. Explain that a large loss with random initial weights is expected and reduces as weights learn.'
    ],
    '7': [
        'Explain that VAEs are useful because they learn a continuous latent space, but their design choices can limit image sharpness and realism.',
        'Discuss the Gaussian posterior and prior assumptions. They simplify training but may not represent all complicated real-data structures exactly.',
        'Explain blur using MSE: when several outputs are plausible, a pixel-wise average may minimise loss even though the average image has soft edges and lost details.',
        'Explain the reconstruction-KL trade-off. Reconstruction wants z to preserve information about x, while KL wants every posterior close to N(0,I). Improving one can harm the other.',
        'Define posterior collapse: the encoder outputs mu near zero and sigma near one for all inputs, so z does not identify the input and KL becomes very small.',
        'Conclude that GANs are motivated by the desire for realistic, detailed samples without explicitly selecting a pixel-wise reconstruction likelihood.'
    ],
    '8': [
        'Introduce the two players: generator G_phi converts noise z into a fake sample; discriminator D_theta returns a probability that its input is real.',
        'Describe generator training: sample z, form G(z), pass it to D, use target y=1, and minimise L_G=-log D(G(z)). The generator wants its fake output judged real.',
        'Explain the backward pass: gradients travel through D to G, but discriminator parameters theta are frozen. The discriminator supplies a signal, not an update target in this step.',
        'Describe discriminator training with real x labelled 1 and generated G(z) labelled 0. It learns L_real=-log D(x) and L_fake=-log[1-D(G(z))].',
        'Add the two parts: L_D=L_real+L_fake. During this discriminator update, generator parameters phi are frozen so only theta changes.',
        'Finish with alternating updates. A balanced competition gradually strengthens both networks: D becomes a better judge and G learns to make more realistic samples.'
    ],
    '9': [
        'Start from binary cross-entropy and derive the discriminator terms. For real x, the target is 1, giving -log D(x). For fake G(z), target 0 gives -log[1-D(G(z))].',
        'State that minimising the negative loss is equivalent to maximising the discriminator value. Use expectations because training averages over many real samples and many noise vectors.',
        'Write V(D,G)=E_x[log D(x)]+E_z[log(1-D(G(z)))]. Explain that D maximises it while G is fixed.',
        'Explain the real term: as D(x) moves from 0.1 to 0.9, log D(x) becomes less negative and larger, so maximisation pushes real predictions toward one.',
        'Explain the fake term: maximising log(1-D(G(z))) pushes D(G(z)) toward zero, so the discriminator calls generated samples fake.',
        'For the generator, minimise the fake term (or use non-saturating -log D(G(z)) in practice). Conclude with min_G max_D V(D,G) and opposite goals.'
    ],
    '10': [
        'Define p_data(x) as the real-data distribution and p_g(x) as the distribution of generator outputs. GAN convergence is about matching these distributions, not copying one image.',
        'At the start, p_g differs strongly from p_data, so generated samples look unrealistic. The discriminator can therefore output near one for real and near zero for fake.',
        'As generator updates use the discriminator gradient, generated samples become more realistic. The value D(G(z)) rises because fake images are harder to reject.',
        'State the ideal condition p_g(x)=p_data(x). If sources have the same distribution, there is no useful evidence for a discriminator to identify their source.',
        'Therefore the optimal discriminator gives 0.5 for real and generated samples. Emphasise that 0.5 here is success, not poor training.',
        'Define Nash equilibrium: neither G nor D can improve its own objective by changing only itself while the other network remains fixed.'
    ],
}

model_answers = {
    '1': 'KL divergence is used when we want to know how different two probability distributions are. Let P(x) be the real distribution and Q(x) be the distribution learned by a model. The forward form, D_KL(P||Q), checks whether the model Q covers all places where real data P occur. If real data have probability but the model gives zero probability there, the penalty is extremely large. Therefore forward KL is mode-covering. In contrast, reverse KL, D_KL(Q||P), strongly penalizes a model that puts probability in places where real data are very unlikely. It therefore chooses a safe, high-density region and is called mode-seeking. For two groups of people around 160 cm and 190 cm, forward KL may cover both but also include 175 cm; reverse KL may model one group well but ignore the other. Thus the direction of KL decides the kind of error that is punished.',
    '2': 'Self-information tells us how surprising one event is. If rain has probability 0.125, it is more surprising than sunshine with probability 0.875, so rain has higher information. Entropy extends this idea to the whole distribution: before we know tomorrow’s weather, it tells us the average uncertainty. Entropy is high when possibilities are equally likely and low when one outcome is almost certain. Cross-entropy asks a different question: how surprised is our model Q by outcomes that actually come from P? In classification, the true class is usually one-hot, so the loss is simply negative log of the probability assigned to the correct class. Finally, KL divergence is cross-entropy minus entropy. It measures the extra cost of using the model distribution Q when the true distribution is P.',
    '3': 'A VAE is a generative model, which means it does not only copy an input; it also learns how to make new similar samples. The input x may be a large image or a feature vector. The encoder compresses x into a smaller latent representation z, but unlike an ordinary autoencoder it gives a Gaussian distribution for z. It outputs a mean mu and a log-variance. We then sample z from that Gaussian. The decoder receives z and reconstructs x-hat. The true posterior p(z|x) is difficult because computing its normalising evidence requires considering all possible z values. Therefore the encoder learns an approximate posterior q_phi(z|x). A standard normal prior N(0,I) gives the latent space a simple, continuous structure, making it possible to generate new data by sampling z and decoding it.',
    '4': 'The decoder is the part of a VAE that changes a latent vector z back into data space. For continuous data, it can output the mean and variance of a Gaussian distribution; for binary data, it outputs a probability for each feature or pixel. We want the decoder to give high probability to the original input x when z has been sampled from the encoder. Directly maximizing log p_theta(x) is hard because it requires integrating over all possible z. The ELBO solves this by using a computable lower bound. Its first term rewards reconstruction: x-hat should be close to x. Its second term is KL regularization: q_phi(z|x) should remain close to the prior p(z). Training minimizes negative ELBO, so VAE loss is reconstruction loss plus KL loss. A good answer must state that both parts are necessary: reconstruction preserves information, while KL makes the latent space useful for sampling.',
    '5': 'In a VAE, the encoder produces mu and sigma, and we must sample a latent vector z. Direct random sampling is a problem for ordinary backpropagation because it hides the dependency between z and the encoder outputs. The reparameterization trick rewrites sampling in a useful way. First, sample epsilon from a standard normal distribution. Then calculate z = mu + sigma times epsilon. The randomness is now entirely inside epsilon, which is independent of the learned encoder parameters. During one backward pass epsilon is treated as a constant. Therefore dz/dmu equals 1 and dz/dsigma equals epsilon. The loss can now send gradients from the decoder through z to mu and sigma, and finally to encoder parameters. The trick does not remove randomness; it moves randomness into a separate input so learning remains differentiable.',
    '6': 'A complete VAE forward pass has a fixed order. Start with the input vector x and multiply it by encoder weights, add bias, and apply ReLU to create hidden features. From those features, use one linear layer to calculate mu and another to calculate log sigma squared. Exponentiate log sigma squared to obtain variance, and take the square root for sigma. Sample epsilon from N(0,I), then obtain z = mu + sigma epsilon. Pass z through the decoder hidden layer and output layer to obtain x-hat. Use a linear output for continuous values and sigmoid for binary or normalized values. Finally calculate reconstruction loss, such as MSE, and calculate the KL divergence between the encoder Gaussian and N(0,I). Add them to obtain total VAE loss. This sequence is important because confusing log-variance, variance, and standard deviation causes numerical errors.',
    '7': 'VAEs learn useful latent spaces, but they can create smooth or blurry images. One reason is pixel-wise reconstruction loss such as MSE. When two sharp outputs are both possible, the average of them may minimise MSE, even though the average looks blurry. A VAE also assumes simple Gaussian forms for the latent posterior and prior, which may not match every real-data pattern. There is another tension: reconstruction loss wants z to keep detailed information about x, while KL loss wants the distribution for every input to look like N(0,I). If KL becomes too strong, posterior collapse can occur. The encoder then gives almost the same N(0,I) distribution for every input, so z no longer describes the input. GANs are motivated by these limitations: they aim for perceptually realistic detail by learning from a discriminator instead of explicit pixel-wise reconstruction loss.',
    '8': 'A GAN contains a generator and a discriminator. The generator receives random noise z and turns it into a fake sample G(z). The discriminator receives either a real sample or a generated sample and outputs a number between zero and one: its estimated probability that the sample is real. Training alternates between two steps. In generator training, the discriminator is frozen. We label G(z) as if it were real, target one, and update only generator parameters so D(G(z)) becomes larger. In discriminator training, the generator is frozen. Real data have label one and generated data have label zero. The discriminator learns to give high scores to real samples and low scores to fake samples. Alternation matters because one network creates the challenge and the other provides the learning signal. Over time, the generator makes more convincing samples and the discriminator becomes a better critic.',
    '9': 'The GAN objective can be built from binary cross-entropy. For a real sample x, the discriminator should output D(x) close to one, giving loss -log D(x). For a generated sample G(z), it should output D(G(z)) close to zero, giving loss -log[1-D(G(z))]. The discriminator minimizes these losses, or equivalently maximizes the value function E[log D(x)] plus E[log(1-D(G(z)))]. The expectations mean that we average across many examples. The generator cannot change the real-data term because it does not contain generator parameters. It changes only the fake term. Its original minimax goal is to make log(1-D(G(z))) small, which pushes D(G(z)) toward one. Thus the complete game is min_G max_D V(D,G). The discriminator wants a fake score near zero; the generator wants the same score near one.',
    '10': 'GAN convergence is described using distributions, not just one attractive generated image. Let p_data be the distribution of real training data and p_g be the distribution produced by the generator. At the beginning, p_g is far from p_data, so fake samples are easy to identify: D(x) is near one for real data and D(G(z)) is near zero for generated data. Through alternating training, the generator changes p_g so it becomes closer to p_data. At ideal convergence, p_g equals p_data. Then a discriminator has no reliable evidence that tells whether a sample is real or generated, so its best output is 0.5 for either source. This does not mean that the discriminator failed. It means the generator has matched the real distribution. This state is a Nash equilibrium because neither network can improve alone while the other stays fixed.'
}

what_it_does = {
    '1': ['P(x) tells where real data occur.', 'Q(x) tells where the model thinks data occur.', 'KL tells how different these two probability maps are.', 'Forward KL says do not miss real groups; reverse KL says do not invent unlikely groups.'],
    '2': ['Self-information means surprise of one result.', 'Entropy means uncertainty before the result is known.', 'Cross-entropy means how surprised the model is by the correct result.', 'KL means the extra error from using the model distribution.'],
    '3': ['Encoder reads x and describes its hidden code.', 'Mu is the centre of possible codes.', 'Sigma is how spread out those codes can be.', 'z is the compact sampled code.', 'Decoder changes z back into data.'],
    '4': ['Decoder likelihood says what output is likely for z.', 'Reconstruction term checks if the output resembles the input.', 'KL term keeps latent codes organised.', 'ELBO and VAE loss combine these two jobs.'],
    '5': ['Epsilon supplies randomness.', 'Mu chooses the centre.', 'Sigma controls how much randomness can move z.', 'The trick keeps z connected to mu and sigma for learning.'],
    '6': ['Encoder hidden layer finds useful patterns.', 'Mu head chooses the typical latent code.', 'Log-variance head chooses uncertainty.', 'Sampling chooses one related code.', 'Decoder rebuilds the data.', 'Loss checks copying quality and latent-space order.'],
    '7': ['Reconstruction loss can encourage averaging.', 'KL loss can remove too much input detail.', 'Posterior collapse means z stops carrying useful input information.', 'GAN adds a realism judge.'],
    '8': ['Noise is the starting material.', 'Generator creates a fake sample.', 'Discriminator judges whether a sample looks real.', 'Training alternates so each network learns from the other.'],
    '9': ['The real-data term rewards correct real decisions.', 'The fake-data term rewards correct fake decisions.', 'The generator changes fakes so they receive a real label.', 'Minimax writes this competition in one formula.'],
    '10': ['p_data is the real pattern.', 'p_g is the generated pattern.', 'Convergence means the patterns match.', 'A 0.5 discriminator score means it cannot tell their source apart.']
}

def what_each_part_does(number):
    doc.add_paragraph('What each part does in simple words', style='Heading 2')
    bullets(what_it_does[number])

def detailed_fifteen_mark_answer(number):
    doc.add_paragraph('15 Mark Answer Points', style='Heading 2')
    add_para('Write these points in order after your introduction and diagram. Together they form a complete long-answer explanation.', 'How to write for 15 marks: ')
    bullets(fifteen_mark_points[number])
    doc.add_paragraph('Simple model answer to learn', style='Heading 2')
    add_para(model_answers[number])

def exam_structure():
    doc.add_paragraph('A reliable 15-mark structure', style='Heading 2')
    flow(['Definition / objective', '        ↓', 'Key components or formula', '        ↓', 'Working or derivation', '        ↓', 'Small numerical or real-world example', '        ↓', 'Advantages, limitation, and conclusion'])
    add_para('Use this order in long answers. Define every symbol the first time it is used, and explain what the formula means in words. This earns method marks even if an arithmetic detail is missed.')

def topic(num, title, definition, flow_lines, explanation, example, mem):
    doc.add_paragraph(f'{num}. {title}', style='Heading 1')
    add_para(definition, 'Core idea: ')
    flow(flow_lines)
    topic_flowchart(num)
    for h, text in explanation:
        doc.add_paragraph(h, style='Heading 2')
        add_para(text)
    doc.add_paragraph('Example', style='Heading 2')
    add_para(example)
    what_each_part_does(num)
    detailed_fifteen_mark_answer(num)
    memory(mem)

add_title('Variational Autoencoders and GANs', 'Complete easy-to-revise notes from Lectures 19 to 24 and 31 to 34')
add_para('This study guide turns the lecture sequence into exam-ready 15-mark answers. It explains each idea in simple language, preserves the lecture notation, and includes short calculations, flowcharts, and memory cues.')
exam_structure()
doc.add_paragraph('Notation used throughout', style='Heading 1')
bullets([
    'x: an observed real data sample; z: a latent/noise vector; x-hat: a reconstruction or generated sample.',
    'p: a true, prior, likelihood, or model distribution; q: an approximate distribution. theta denotes decoder/discriminator parameters and phi denotes encoder/generator parameters.',
    'E[ ] means expectation (average over samples). N(mu, sigma squared) is a Gaussian distribution. I is the identity matrix.',
    'For information quantities, the lecture uses log base 2, so values are in bits. GAN loss examples use natural log, which is standard in optimization.'
])

topic('1', 'KL Divergence and Its Direction',
      'KL divergence measures how different an approximate probability distribution Q is from a reference distribution P. It is not an ordinary distance: it is directional, so D_KL(P || Q) is generally not equal to D_KL(Q || P).',
      ['True distribution P(x)', '        ↓ compare', 'Learned distribution Q(x)', '        ↓', 'KL divergence: mismatch penalty'],
      [('Formula', 'For discrete variables, D_KL(P || Q) = sum_x P(x) log[P(x)/Q(x)]. For continuous variables, replace the sum by an integral. KL is zero only when the distributions match almost everywhere.'),
       ('Forward and reverse KL', 'Forward KL D_KL(P || Q) heavily penalizes Q if Q(x)=0 where P(x)>0. Hence it tries to cover all regions where real data occur. Reverse KL D_KL(Q || P) heavily penalizes Q for placing probability where P is tiny, so it concentrates on high-density regions.'),
       ('Mode behaviour', 'With two separated real-data modes, forward KL is mode-covering: a simple single Gaussian may spread across both modes and also cover the low-density gap. Reverse KL is mode-seeking: it may select one mode, generate sharp samples there, and miss the other valid mode.')],
      'Heights of people may cluster near 160 cm and 190 cm with few people near 175 cm. A limited forward-KL model may broaden to cover both groups and place some probability near 175 cm. A reverse-KL model may concentrate on one cluster, giving realistic samples but missing the other.',
      ['Write the direction before interpreting it: P || Q means “does Q cover P?”', 'Forward KL: missing a real-data region is very costly; remember “cover.”', 'Reverse KL: placing mass in an unrealistic region is costly; remember “seek.”', 'State clearly that KL is non-negative and asymmetric.'])
visual('lec19-09.png', 'Forward KL and reverse KL on a two-cluster height distribution.', 'The blue curve is the real data with two groups. In the left graph, the orange model spreads wide to cover both groups, but it wrongly puts some probability in the middle gap. In the right graph, the green model fits one group closely, but misses the other group. In an exam, label these as mode-covering and mode-seeking.')

topic('2', 'Self Information Entropy Cross Entropy and KL',
      'Self-information measures surprise for one observed outcome. Entropy measures average uncertainty before the outcome is known. Cross-entropy measures the surprise assigned by a model to the true outcomes.',
      ['Probability of outcome', '        ↓', 'Self-information: -log P(x)', '        ↓ average using P(x)', 'Entropy: H(P)', '        ↓ use model Q instead', 'Cross-entropy: H(P,Q)', '        ↓ subtract H(P)', 'KL: D_KL(P || Q)'],
      [('Self-information', 'I(x) = -log2 P(x) = log2[1/P(x)]. A likely event gives little surprise; an unlikely event gives high surprise. The logarithm makes information from independent events add.'),
       ('Entropy', 'H(P) = -sum_x P(x) log2 P(x). It is the probability-weighted average surprise of all possible outcomes. A fair two-outcome distribution has H=1 bit; a certain outcome has H=0.'),
       ('Cross-entropy and derivation', 'H(P,Q) = -sum_x P(x) log2 Q(x). Thus D_KL(P || Q) = H(P,Q) - H(P) = sum_x P(x) log2[P(x)/Q(x)]. When P is one-hot for the correct class, cross-entropy becomes -log Q(correct class).')],
       'Weather is an easy way to remember the ideas. A highly likely sunny day has low surprise; an unlikely rainy day has high surprise. Entropy describes uncertainty before tomorrow arrives. Cross-entropy becomes small when the model gives high probability to the weather that actually occurs and large when it gives the correct weather low probability.',
      ['One event: self-information; all possible events: entropy.', 'Entropy uses the true distribution P in both the weight and log term.', 'Cross-entropy uses P to weight outcomes but Q inside the log.', 'KL equals extra coding/surprise cost: cross-entropy minus entropy.'])

topic('3', 'VAE Architecture and Encoder',
      'A variational autoencoder is an unsupervised generative model. It learns a structured latent representation so that it can reconstruct inputs and generate new samples.',
      ['Input x', '  ↓ Encoder q_phi(z|x)', 'mean mu(x), log-variance log sigma^2(x)', '  ↓ sample z', 'latent vector z', '  ↓ Decoder p_theta(x|z)', 'reconstruction x-hat'],
      [('Why a VAE is needed', 'An ordinary autoencoder learns reconstruction. A VAE additionally learns a distribution in latent space, enabling samples z to be decoded into new data. It aims to learn p_theta(x) close to the unknown p_data(x).'),
       ('Two mappings', 'The encoder approximates the intractable true posterior p(z|x) with q_phi(z|x). The decoder represents the likelihood p_theta(x|z). Bayes rule gives p(z|x)=p(x|z)p(z)/p(x), but p(x)=integral p(x|z)p(z) dz is hard to compute.'),
       ('Gaussian latent space', 'The usual prior is p(z)=N(0,I). The encoder outputs a mean vector mu_phi(x) and log-variance vector log sigma_phi^2(x), defining q_phi(z|x)=N(mu_phi(x), sigma_phi^2(x)). Log-variance is output for numerical convenience; variance is exp(log-variance).')],
       'Think of every latent dimension as one hidden feature. The encoder gives two descriptions for each feature: where it is centred (mean) and how uncertain it is (variance). Sampling takes one reasonable value from every hidden feature and joins them into z.',
      ['Data space is x; latent space is z. They need not have the same dimension.', 'Encoder output is a distribution, not one fixed code.', 'Remember q for approximate posterior and p for prior/likelihood.', 'The KL term later makes q_phi(z|x) close to N(0,I).'])
visual('lec21-06.png', 'VAE map from data space to latent space and back.', 'Read the arrows from left to right. Start with an input x. The encoder gives a probability distribution for z, not one fixed z. Then the decoder uses z to rebuild x-hat. The central box p(z) is the simple prior pattern the latent space is encouraged to follow.')

topic('4', 'VAE Decoder ELBO and Loss',
      'The decoder turns a sampled latent vector into parameters of a distribution over data. VAE training maximizes the evidence lower bound (ELBO), or equivalently minimizes reconstruction loss plus KL regularization.',
      ['x → encoder → q_phi(z|x)', '        ↓ sample z', 'decoder gives p_theta(x|z)', '        ↓', 'reconstruct x-hat and compute loss', '        ↓', 'reconstruction term + KL term'],
      [('Decoder output depends on data type', 'For continuous data the decoder can output mu_theta(z) and sigma_theta squared(z) of a Gaussian p_theta(x|z). For binary data or normalized pixels it outputs Bernoulli probabilities pi_theta(z). Different z values sampled for the same x can produce slightly different but similar reconstructions.'),
       ('ELBO', 'ELBO(x) = E_q[log p_theta(x|z)] - D_KL(q_phi(z|x) || p(z)). It is a computable lower bound on log p_theta(x). We maximize it because direct evaluation of the evidence p_theta(x) requires integration over all z.'),
       ('Loss', 'Minimizing negative ELBO gives L_VAE = -E_q[log p_theta(x|z)] + D_KL(q_phi(z|x)||p(z)). The first term rewards accurate reconstruction. Use MSE for continuous outputs and BCE for binary outputs. For one latent dimension versus N(0,1), KL = 1/2[mu squared + sigma squared - log(sigma squared) - 1]; sum this over dimensions.')],
       'For continuous inputs, MSE simply asks “How far is the reconstruction from the original?” The KL term asks “Does this input’s latent distribution still look like the standard normal prior?” A VAE learns best when it can answer both questions well.',
      ['ELBO has two ideas: reconstruct well and keep latent space organized.', 'Loss is negative ELBO, so the signs reverse when training by minimization.', 'Choose MSE for continuous data, BCE for binary data.', 'State why ELBO is needed: the evidence integral is intractable.'])
visual('lec22-07.png', 'The ELBO split into reconstruction and KL regularization.', 'The blue-marked term asks: “Can the decoder rebuild the input?” The red-marked term asks: “Is the encoder distribution close to the simple prior?” For loss minimization, turn the first term negative and add the KL term. Remember the short form: reconstruct well plus keep z organized.')

topic('5', 'Reparameterization Trick',
      'Sampling directly from q_phi(z|x) blocks ordinary backpropagation through the random operation. The reparameterization trick keeps randomness separate from the learnable encoder parameters.',
      ['Encoder produces mu and sigma', '        ↓', 'sample epsilon ~ N(0,I)', '        ↓', 'z = mu + sigma elementwise-multiply epsilon', '        ↓', 'decoder → loss → gradients to encoder'],
      [('Key equation', 'Instead of writing z sampled from N(mu,sigma squared), write z = mu_phi(x) + sigma_phi(x) elementwise-multiply epsilon, where epsilon sampled from N(0,I). epsilon is random but independent of phi, and is treated as constant during a backward pass.'),
       ('Why gradients now flow', 'The decoder is differentiable, so gradients already reach z. With the new expression, partial z/partial mu = 1 and partial z/partial sigma = epsilon. Therefore gradients can continue from z to mu, sigma, and then to encoder parameters phi.'),
       ('End-to-end update', 'By chain rule, partial L/partial phi = (partial L/partial mu)(partial mu/partial phi) + (partial L/partial sigma)(partial sigma/partial phi). The stochastic part is isolated in epsilon; the rest is deterministic for that sample.')],
       'Imagine mu as the starting location, sigma as the size of the allowed movement, and epsilon as a random direction. The formula takes the starting location and moves it by a random amount. Because the formula shows the movement clearly, the network can learn how to adjust the start and spread.',
      ['Do not say sampling disappears; randomness remains in epsilon.', 'What changes is the path: z becomes differentiable with respect to mu and sigma.', 'Know the two derivatives: dz/dmu=1 and dz/dsigma=epsilon.', 'Use the phrase “separates randomness from learnable parameters.”'])
visual('lec23-07.png', 'Reparameterization trick and its two key derivatives.', 'The red equation is the only line to memorise: z = mu + sigma times epsilon. Think of epsilon as a random dice roll that the network does not learn. The encoder learns only mu and sigma, so gradients can pass through them. The boxed results explain why: changing mu shifts z directly, while changing sigma scales the random value epsilon.')

topic('6', 'Complete VAE Working in Simple Words',
      'A full VAE pass has two simple jobs: understand an input as a compact uncertain code, then use that code to rebuild the input while learning from its mistakes.',
      ['x → encoder hidden layer + ReLU', '        ↓', 'mu and log sigma^2 (linear heads)', '        ↓', 'sigma = sqrt(exp(log sigma^2))', '        ↓', 'z = mu + sigma epsilon', '        ↓', 'decoder hidden layer + ReLU → linear/sigmoid output', '        ↓', 'MSE or BCE + KL'],
       [('Step 1: Understand the input', 'The encoder sees the input and produces hidden features. These are useful patterns it has noticed, such as a general shape or an important combination of features.'),
        ('Step 2: Describe uncertainty', 'The encoder then gives mu and log-variance. Mu means the most typical hidden code for this input. Variance means how much flexibility there is around that typical code.'),
        ('Step 3: Rebuild and learn', 'The model samples z, lets the decoder rebuild x-hat, and checks whether the rebuild resembles the input and whether the latent code stays organised. Training improves both together.')],
       'Imagine writing a compact note for a friend. The encoder writes the note, sampling chooses one related version of it, and the decoder uses it to redraw the original. The loss tells the model whether the redraw is clear and whether the notes are kept in an organised filing system.',
      ['Mean and log-variance heads normally use no ReLU: their outputs may be negative.', 'Convert log sigma squared with exp; take square root to obtain sigma.', 'Use sampled z, not only mu, when tracing the stochastic VAE forward pass.', 'Final loss = reconstruction loss + KL divergence loss.'])

topic('7', 'Limitations of VAEs and Motivation for GANs',
      'VAEs are valuable for structured latent variables and reconstruction, but their likelihood assumptions and pixel-wise losses can make generated images smooth or blurry. GANs were introduced to learn realistic samples through competition rather than explicit reconstruction likelihood.',
      ['VAE strengths: latent structure and generation', '        ↓ limitations', 'Gaussian assumptions / blur / trade-off / collapse', '        ↓ motivation', 'GAN: adversarial realism'],
      [('Main limitations', 'A standard VAE commonly assumes Gaussian q_phi(z|x) and p(z)=N(0,I), which may be too simple for complex data. MSE can average several plausible outputs, removing sharp edges and fine texture. The reconstruction term wants z to retain input-specific detail, whereas the KL term pushes q_phi(z|x) toward the prior.'),
       ('Posterior collapse', 'If KL pressure is too strong, the encoder can output mu near 0 and sigma near 1 for every input. Then q_phi(z|x) is nearly N(0,I) for all x, KL is near zero, and z no longer carries useful information about a particular input.'),
       ('Likelihood assumption', 'The decoder explicitly assumes a data likelihood: often Gaussian/MSE for continuous data or Bernoulli/BCE for binary data. If this choice does not fit the data, uncertainty and fine details may be represented poorly.')],
      'For a face image, pixel-wise MSE may average two plausible edge locations. The average has a soft transition rather than a sharp edge, so the output can look blurry even if its MSE is low.',
      ['List the limitation, then explain the mechanism, then state its effect.', 'Posterior collapse means latent z becomes uninformative, not merely “KL is small.”', 'The VAE reconstruction-KL balance is a genuine trade-off.', 'GAN motivation: improve perceptual realism without choosing a pixel-wise reconstruction likelihood.'])

topic('8', 'GAN Architecture and Alternating Training',
      'A generative adversarial network has two neural networks: generator G_phi transforms noise z into a fake sample; discriminator D_theta outputs the probability that its input is real. They improve through alternating updates.',
      ['noise z ~ p(z)', '        ↓', 'Generator G_phi(z) = x_fake', '        ↓', 'Discriminator D_theta(x_fake)', '        ↓', 'probability input is real'],
      [('Generator update', 'Freeze theta. Sample z, compute x_fake=G_phi(z), pass it through D_theta, and set the desired label to 1. The commonly used non-saturating generator loss is L_G=-log D_theta(G_phi(z)). Gradients pass through the fixed discriminator to update only phi.'),
       ('Discriminator update', 'Freeze phi. Use real x with target 1 and generated G_phi(z) with target 0. L_real=-log D_theta(x). L_fake=-log[1-D_theta(G_phi(z))]. Minimize L_D=L_real+L_fake with respect to theta.'),
       ('Why freeze one network', 'During a generator step, the discriminator provides a differentiable learning signal but must not change. During a discriminator step, the generator simply supplies fake samples and must not change. This alternating optimization is essential.')],
      'If D(G(z))=0.2 during a generator update, L_G=-ln(0.2) about 1.609, so the generator receives a strong signal to make future outputs more real-looking. If D(x)=0.9 and D(G(z))=0.2 during a discriminator update, L_D=-ln(0.9)-ln(0.8) about 0.328.',
      ['Generator target is 1 when training G; fake target is 0 when training D.', 'Freeze D for G update; freeze G for D update.', 'D output means probability of “real,” not probability of “fake.”', 'Generator creates; discriminator judges.'])
visual('lec32-03.png', 'Generator-training graph in a GAN.', 'Follow the green arrows. Noise z enters the generator, which creates a fake image. The discriminator judges that image. During this step, the lock means the discriminator is frozen: gradients travel through it, but only the generator parameters change. The generator wants the judge to output 1, meaning “real.”')

topic('9', 'GAN Objectives and Minimax Game',
      'The discriminator maximizes its ability to separate real and fake samples. The generator minimizes the same value function so that generated samples are classified as real.',
      ['Real x ~ p_data → maximize log D(x)', 'Fake: z ~ p_z → G(z) → maximize log[1-D(G(z))] for D', '                                 ↓', 'G instead minimizes log[1-D(G(z))]', '                                 ↓', 'min_G max_D V(D,G)'],
      [('Discriminator objective', 'V_D(theta;phi) = E_{x~p_data}[log D_theta(x)] + E_{z~p_z}[log(1-D_theta(G_phi(z)))]. Maximize it with respect to theta while holding phi fixed. Higher D(x) raises the first term; lower D(G(z)) raises the second term.'),
       ('Generator objective', 'In the original minimax formulation, G minimizes E_z[log(1-D(G(z)))] while D is fixed. This pushes D(G(z)) toward 1. In practical training, the non-saturating loss -E_z[log D(G(z))] is often used because it gives stronger gradients early in training; it has the same desired direction.'),
       ('Combined game', 'The original GAN objective is min_G max_D E_x[log D(x)] + E_z[log(1-D(G(z)))]. The two networks have opposing objectives on generated samples: D wants them near 0; G wants them near 1.')],
      'For a real sample, D(x)=0.1 gives log(0.1)=-2.303, whereas D(x)=0.9 gives log(0.9)=-0.105, a larger value closer to zero. Thus maximizing log D(x) pushes real predictions toward 1. For a fake sample, maximizing log(1-D(G(z))) pushes D(G(z)) toward 0.',
      ['Define both expectations and say which parameters are fixed in each update.', 'Objective maximization equals minimizing the negative loss.', 'Original minimax generator loss and non-saturating generator loss are different expressions but share the goal D(G(z))→1.', 'End with min_G max_D to show the adversarial relationship.'])

topic('10', 'GAN Convergence and Nash Equilibrium',
      'Ideal GAN convergence occurs when the generator distribution matches the real-data distribution. At that point the best discriminator cannot distinguish source and outputs probability 0.5 for either source.',
      ['Early: p_g differs from p_data', '        ↓', 'D(x) near 1; D(G(z)) near 0', '        ↓ alternating updates', 'p_g moves toward p_data', '        ↓ ideal equilibrium', 'p_g = p_data and D(x)=D(G(z))=0.5'],
      [('Training progression', 'Initially generated samples may look like noise, so D separates them easily. As the generator receives gradients through the discriminator, p_g(x) moves closer to p_data(x), and D(G(z)) may progress from 0.1 to 0.4 to 0.5.'),
       ('Ideal convergence', 'When p_g(x)=p_data(x), real and generated samples have the same distribution. A discriminator has no information about their source, so its optimal output is 1/2 for either kind of sample. A 0.5 output here means successful distribution matching, not a badly trained discriminator.'),
       ('Nash equilibrium', 'A Nash equilibrium is a state where neither player can improve its own outcome by changing only its own parameters while the other player remains fixed. In the ideal GAN state, D cannot improve because the distributions match, and G cannot improve its match further.')],
      'For faces, early generated samples may be distorted and D(G(z)) close to 0. With effective training, generated faces become harder to classify. At the ideal theoretical endpoint, a real face and a generated face are each assigned probability 0.5 of being real.',
      ['Write p_g=p_data as the convergence condition.', 'Write D=0.5 as the consequence, not as a failure.', 'Nash equilibrium requires neither player to improve unilaterally.', 'Practical GAN training may be unstable; this describes the ideal theoretical result.'])
visual('lec34-5.png', 'Ideal GAN convergence.', 'This graph gives the final exam statement: p_g(x) = p_data(x). When generated and real samples follow the same distribution, the discriminator has no clue where a sample came from. It therefore outputs 0.5 for both. This is a success, not a weak discriminator.')

doc.add_paragraph('VAE and GAN Comparison', style='Heading 1')
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'
headers = ['Aspect', 'VAE', 'GAN']
for cell, value in zip(table.rows[0].cells, headers):
    cell.text = value; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), '1F4E78'); tcPr.append(shd)
    for r in cell.paragraphs[0].runs: r.font.color.rgb = RGBColor(255,255,255); r.bold = True
for row in [
    ('Core idea', 'Probabilistic encoder-decoder with a regularized latent space.', 'Generator and discriminator compete in an adversarial game.'),
    ('Training signal', 'Reconstruction likelihood plus KL divergence.', 'Discriminator feedback through BCE/minimax objectives.'),
    ('Strength', 'Structured, sampleable latent representation.', 'Often sharp, realistic samples.'),
    ('Common concern', 'Blurry samples or posterior collapse.', 'Training instability; balance between G and D.')
]:
    cells = table.add_row().cells
    for c, value in zip(cells, row):
        c.text = value; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs: p.paragraph_format.space_after = Pt(2)

doc.add_paragraph('Final Revision Checklist', style='Heading 1')
bullets([
    'Can you write and explain D_KL(P || Q), entropy, cross-entropy, ELBO, VAE loss, reparameterization, GAN losses, and the minimax objective?',
    'Can you label which network is frozen during each GAN update?',
    'Can you distinguish variance, log-variance, and standard deviation in a VAE numerical problem?',
    'Can you explain why ideal GAN discriminator output 0.5 is success?',
    'For every 15-mark response: definition → formula → working flowchart → example → interpretation → limitation/conclusion.'
])

doc.add_paragraph('One Minute Memory Stories', style='Heading 1')
add_para('Use these short stories when the formulas feel difficult. First remember the story; then attach the formula to it.')
for title, text in [
    ('KL divergence', 'Imagine P is the map of where people really live and Q is your model map. Forward KL punishes you if your map leaves out a real town. Reverse KL punishes you if your map invents a town where nobody lives.'),
    ('Entropy and cross-entropy', 'Entropy asks “How uncertain is the weather before tomorrow?” Cross-entropy asks “How surprised was my model by the correct weather?” A correct answer with high predicted probability gives a small loss.'),
    ('VAE encoder', 'The encoder does not put x in one fixed drawer. It gives a cloud location: centre mu and spread sigma. Sampling chooses one point from that cloud.'),
    ('VAE loss', 'A VAE student has two jobs: copy the input clearly and keep its latent notebook neat. Reconstruction checks the copy; KL checks the neat notebook.'),
    ('Reparameterization', 'Instead of hiding randomness inside the encoder, take the random dice roll epsilon outside. Then the network can still see how mu and sigma change z.'),
    ('GAN', 'The generator is a counterfeiter; the discriminator is an inspector. The counterfeiter learns from each inspection. Train one while holding the other still.'),
    ('GAN equilibrium', 'When the counterfeits are indistinguishable from genuine items, the inspector can do no better than a 50-50 guess. That is the ideal equilibrium.')
]:
    doc.add_paragraph(title, style='Heading 2')
    add_para(text)

doc.core_properties.title = 'Variational Autoencoders and GANs Study Guide'
doc.core_properties.subject = 'Exam revision notes based on Lectures 19 to 24 and 31 to 34'
doc.save(OUT)
print(OUT.resolve())
